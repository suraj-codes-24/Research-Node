from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

# Create a new presentation
prs = Presentation()

# Define slide layouts
# 0: Title, 1: Title and Content, 2: Section Header, 3: Two Content, 5: Title Only, 6: Blank
TITLE_SLIDE_LAYOUT = prs.slide_layouts[0]
CONTENT_SLIDE_LAYOUT = prs.slide_layouts[1]
TWO_CONTENT_LAYOUT = prs.slide_layouts[3]

# Helper to add speaker notes
def add_notes(slide, text):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = text

# --- Slide 1: Title ---
slide = prs.slides.add_slide(TITLE_SLIDE_LAYOUT)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "ResearchNode"
title.text_frame.paragraphs[0].font.size = Pt(54)
title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 102, 204) # Blue

subtitle.text = "A GraphRAG Reasoning Engine for Scientific Literature\n\nTeam: [Your Names Here]\nRoll No: [Your Roll Numbers]"
add_notes(slide, "Good morning respected teachers and judges. We are Team [Name], and today we are excited to present our mini-project: ResearchNode. It is an intelligent reasoning engine designed to fundamentally change how researchers interact with scientific literature using a cutting-edge technique called GraphRAG.")

# --- Slide 2: The Problem Statement ---
slide = prs.slides.add_slide(CONTENT_SLIDE_LAYOUT)
title = slide.shapes.title
title.text = "The Problem Statement"
content = slide.placeholders[1]
tf = content.text_frame
tf.text = "Information Overload: Thousands of papers are published daily."
p = tf.add_paragraph()
p.text = "Disconnected Knowledge: Search engines like Google Scholar find papers but don't connect ideas."
p = tf.add_paragraph()
p.text = "LLM Hallucinations: Standard AI (like ChatGPT) makes up facts and lacks real citations."
add_notes(slide, "The core problem we are solving is information overload in academia. Today, researchers spend countless hours manually cross-referencing papers. While search engines like Google Scholar can find PDFs, they cannot tell you how Paper A's methodology relates to Paper B's dataset. Furthermore, if you ask standard AI like ChatGPT, it often hallucinates facts and fails to provide verifiable citations.")

# --- Slide 3: Our Solution: ResearchNode ---
slide = prs.slides.add_slide(CONTENT_SLIDE_LAYOUT)
title = slide.shapes.title
title.text = "Our Solution: ResearchNode"
content = slide.placeholders[1]
tf = content.text_frame
tf.text = "Upload: Automatically extract text and entities from research PDFs."
p = tf.add_paragraph()
p.text = "Connect: Build a visual, persistent Knowledge Graph of science."
p = tf.add_paragraph()
p.text = "Reason: Ask complex questions and get answers using GraphRAG."
add_notes(slide, "Our solution is ResearchNode. It does three things: First, you upload research PDFs, and our system extracts the text and identifies key entities. Second, it connects these entities into a persistent Knowledge Graph, allowing you to literally see the web of research. Third, it allows you to reason over this data. You can ask complex questions in plain English, and our GraphRAG engine gives you answers backed by exact citations.")

# --- Slide 4: System Architecture ---
slide = prs.slides.add_slide(CONTENT_SLIDE_LAYOUT)
title = slide.shapes.title
title.text = "System Architecture"
content = slide.placeholders[1]
tf = content.text_frame
tf.text = "1. Processing Pipeline:"
p = tf.add_paragraph()
p.text = "Text Chunking → MiniLM Embeddings → Qdrant (Vector DB)"
p.level = 1
p = tf.add_paragraph()
p.text = "Entity Extraction (spaCy + Gemini) → Neo4j (Graph DB)"
p.level = 1
p = tf.add_paragraph()
p.text = "2. Query Pipeline (GraphRAG):"
p = tf.add_paragraph()
p.text = "User Question → Hybrid Retriever (Vector + Graph) → Gemini LLM → Answer"
p.level = 1
add_notes(slide, "Here is the high-level architecture. When a PDF is uploaded, it takes two parallel paths. Path 1: The text is chunked, embedded using the all-MiniLM model, and stored in Qdrant for semantic search. Path 2: We use spaCy and Gemini to extract scientific entities, storing them in a Neo4j Graph Database. When a user asks a question, our Hybrid Retriever pulls context from both databases simultaneously to generate the final answer.")

# --- Slide 5: What is GraphRAG? ---
slide = prs.slides.add_slide(TWO_CONTENT_LAYOUT)
title = slide.shapes.title
title.text = "Why GraphRAG? (Vs. Standard RAG)"
left_content = slide.placeholders[1].text_frame
right_content = slide.placeholders[2].text_frame

left_content.text = "Standard RAG"
p = left_content.add_paragraph()
p.text = "Uses only Vector Search."
p.level = 1
p = left_content.add_paragraph()
p.text = "Finds text that sounds similar."
p.level = 1
p = left_content.add_paragraph()
p.text = "Misses nuanced relationships."
p.level = 1

right_content.text = "GraphRAG"
p = right_content.add_paragraph()
p.text = "Uses Vector Search + Graph Traversal."
p.level = 1
p = right_content.add_paragraph()
p.text = "Finds similar text AND related concepts."
p.level = 1
p = right_content.add_paragraph()
p.text = "Provides much richer, contextual answers."
p.level = 1
add_notes(slide, "A key innovation in our project is the use of GraphRAG rather than standard RAG. Standard RAG relies only on vector search, which finds text that sounds similar to your question. GraphRAG adds a knowledge graph. So if you ask about 'BERT', our system doesn't just find text mentioning BERT; it traverses the graph to see that BERT is connected to the 'Transformer' model, pulling all that rich context into the AI's prompt for a much deeper answer.")

# --- Slide 6: Multi-Agent System ---
slide = prs.slides.add_slide(CONTENT_SLIDE_LAYOUT)
title = slide.shapes.title
title.text = "Autonomous Multi-Agent System"
content = slide.placeholders[1]
tf = content.text_frame
tf.text = "Literature Discovery Agent:"
p = tf.add_paragraph()
p.text = "Synthesizes a cohesive reading order based on a topic."
p.level = 1
p = tf.add_paragraph()
p.text = "Contradiction Detection Agent:"
p = tf.add_paragraph()
p.text = "Globally scans the graph for papers with conflicting claims."
p.level = 1
p = tf.add_paragraph()
p.text = "Experiment Suggestion Agent:"
p = tf.add_paragraph()
p.text = "Brainstorms novel experiments based on graph gaps and limitations."
p.level = 1
add_notes(slide, "Beyond just answering questions, we implemented a Multi-Agent system to act as autonomous research assistants. The Literature Agent takes a topic and synthesizes a reading order. The Contradiction Agent globally scans the graph to find papers reaching conflicting conclusions. Finally, the Experiment Agent analyzes the limitations sections of your papers and cross-references them with known models in the graph to brainstorm entirely novel research experiments.")

# --- Slide 7: Technology Stack ---
slide = prs.slides.add_slide(TWO_CONTENT_LAYOUT)
title = slide.shapes.title
title.text = "Technology Stack"
left_content = slide.placeholders[1].text_frame
right_content = slide.placeholders[2].text_frame

left_content.text = "Backend & Data"
p = left_content.add_paragraph()
p.text = "Python (FastAPI)"
p.level = 1
p = left_content.add_paragraph()
p.text = "Qdrant (Vector Database)"
p.level = 1
p = left_content.add_paragraph()
p.text = "Neo4j (Knowledge Graph DB)"
p.level = 1

right_content.text = "Frontend & Intelligence"
p = right_content.add_paragraph()
p.text = "React + Vite + Vanilla CSS"
p.level = 1
p = right_content.add_paragraph()
p.text = "Cytoscape.js (Graph Vis)"
p.level = 1
p = right_content.add_paragraph()
p.text = "Google Gemini 1.5 Pro API"
p.level = 1
add_notes(slide, "For our tech stack, we built a modern Single Page Application using React and Vite, utilizing Cytoscape.js to render the interactive knowledge graph. The backend is powered by Python and FastAPI. We use Qdrant for our Vector Database and Neo4j for our Graph Database. The intelligence layer is powered by the Google Gemini API, spaCy for offline Named Entity Recognition, and HuggingFace Sentence Transformers for embeddings.")

# --- Slide 8: Live Demonstration ---
slide = prs.slides.add_slide(TITLE_SLIDE_LAYOUT)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Live Demonstration"
subtitle.text = "Upload → Graph → Chat → Agents"
add_notes(slide, "We will now show a brief live demonstration of the system. (Show dashboard, upload a paper, view graph, ask a question, and test an agent).")

# --- Slide 9: Conclusion & Future Scope ---
slide = prs.slides.add_slide(CONTENT_SLIDE_LAYOUT)
title = slide.shapes.title
title.text = "Conclusion & Future Scope"
content = slide.placeholders[1]
tf = content.text_frame
tf.text = "Conclusion:"
p = tf.add_paragraph()
p.text = "Combining Vector + Graph retrieval drastically improves AI reasoning."
p.level = 1
p = tf.add_paragraph()
p.text = "Future Scope:"
p = tf.add_paragraph()
p.text = "Cloud deployment (AWS/GCP)."
p.level = 1
p = tf.add_paragraph()
p.text = "Real-time integration with arXiv and PubMed APIs."
p.level = 1
p = tf.add_paragraph()
p.text = "Multi-user collaborative workspaces."
p.level = 1
add_notes(slide, "To conclude, ResearchNode successfully demonstrates that combining semantic vector search with explicit graph relationships significantly improves AI reasoning over scientific literature. For future scope, we plan to deploy the system to the cloud, integrate directly with APIs like arXiv for real-time paper ingestion, and add multi-user collaborative workspaces.")

# --- Slide 10: Q&A ---
slide = prs.slides.add_slide(TITLE_SLIDE_LAYOUT)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Thank You!"
subtitle.text = "Any Questions?"
add_notes(slide, "Thank you for your time. We would now be happy to answer any questions.")

# Save presentation
prs.save('ResearchNode_Presentation.pptx')
print("Presentation successfully saved as ResearchNode_Presentation.pptx")
